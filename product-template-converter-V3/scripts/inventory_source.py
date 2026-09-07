from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from lxml import etree
from PIL import Image

from image_fingerprint import canonical_visual_sha256

from pipeline_common import finish, normalize_text, result, sha256_file, stable_id, write_json
from powershell_host_v3 import WINRT_OCR_POWERSHELL_HOST_ENV, resolve_powershell_host


OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def local_name(element: etree._Element) -> str:
    return etree.QName(element).localname


def _ancestors(element: etree._Element):
    parent = element.getparent()
    while parent is not None:
        yield parent
        parent = parent.getparent()


class ItemCollector:
    """Keep inventory order stable and prevent the same OOXML object twice."""

    def __init__(self) -> None:
        self.items: list[dict[str, Any]] = []
        self._source_ids: set[str] = set()

    def add(self, item: dict[str, Any]) -> bool:
        source_id = str(item.get("source_id") or item.get("id") or "")
        if not source_id or source_id in self._source_ids:
            return False
        item["id"] = source_id
        item["source_id"] = source_id
        self._source_ids.add(source_id)
        self.items.append(item)
        return True


def text_item(kind: str, location: str, text: str, **extra: Any) -> dict[str, Any]:
    text = normalize_text(text)
    return {"id": stable_id(kind, location, text), "kind": kind, "location": location, "text": text, **extra}


def package_part_from_rels(name: str) -> str:
    path = Path(name)
    if path.parent.name == "_rels" and path.name.endswith(".rels"):
        return str(path.parent.parent / path.name[:-5]).replace("\\", "/")
    return name


def is_external_target(target: str, target_mode: str) -> bool:
    if target_mode.lower() == "external":
        return True
    parsed = urlparse(target)
    return bool(parsed.scheme and (parsed.netloc or parsed.scheme in {"mailto", "file"}))


def add_external_relationship_items(package: zipfile.ZipFile, collector: ItemCollector, prefix: str = "") -> int:
    added = 0
    for name in sorted(package.namelist()):
        if not name.endswith(".rels"):
            continue
        try:
            root = etree.fromstring(package.read(name))
        except (KeyError, etree.XMLSyntaxError):
            continue
        source_part = package_part_from_rels(name)
        for relationship in root:
            if local_name(relationship) != "Relationship":
                continue
            target = str(relationship.get("Target") or "")
            target_mode = str(relationship.get("TargetMode") or "")
            if not target or not is_external_target(target, target_mode):
                continue
            location = f"{prefix}part:{source_part}/relationship:{relationship.get('Id', '')}"
            details = {
                "source_part": source_part,
                "relationship_id": str(relationship.get("Id") or ""),
                "relationship_type": str(relationship.get("Type") or ""),
                "target": target,
                "target_mode": target_mode or "External",
            }
            collector.add({
                "id": stable_id("relationship", location, json.dumps(details, sort_keys=True)),
                "kind": "relationship",
                "location": location,
                **details,
            })
            added += 1
    return added


def add_ooxml_descriptions(root: etree._Element, part: str, collector: ItemCollector, prefix: str = "") -> int:
    added = 0
    description_index = 0
    for element in root.iter():
        if local_name(element) not in {"cNvPr", "docPr"}:
            continue
        values = {
            "name": str(element.get("name") or ""),
            "description": str(element.get("descr") or element.get("description") or ""),
            "title": str(element.get("title") or ""),
        }
        if not any(normalize_text(value) for value in values.values()):
            continue
        description_index += 1
        location = f"{prefix}part:{part}/description:{description_index}"
        visible_text = normalize_text(" | ".join(value for value in values.values() if normalize_text(value)))
        collector.add({
            "id": stable_id("shape_description", location, visible_text),
            "kind": "shape_description",
            "location": location,
            "text": visible_text,
            **values,
        })
        added += 1
    return added


def image_visual_sha256(data: bytes) -> str:
    return canonical_visual_sha256(data)


def media_item(kind: str, location: str, name: str, data: bytes, **extra: Any) -> dict[str, Any]:
    digest = hashlib.sha256(data).hexdigest().upper()
    source_id = stable_id(kind, location, digest=digest)
    item = {"id": source_id, "source_id": source_id, "kind": kind, "location": location, "name": name, "sha256": digest, "bytes": len(data), **extra}
    if kind == "image":
        item["visual_sha256"] = image_visual_sha256(data)
    return item


def add_pdf_text_item(items: list[dict[str, Any]], page_index: int, line_index: int, text: str, extraction: str = "pdf_text") -> bool:
    text = normalize_text(text)
    if not text:
        return False
    location = f"page:{page_index}/line:{line_index}" if extraction == "pdf_text" else f"page:{page_index}/ocr"
    source_id = stable_id("text", location, text)
    items.append({"id": source_id, "source_id": source_id, "kind": "text", "location": location, "text": text, "extraction": extraction})
    return True


def add_pdf_ocr_text_item(items: list[dict[str, Any]], page_index: int, text: str) -> bool:
    return add_pdf_text_item(items, page_index, 1, text, extraction="ocr")


def locate_pdftoppm() -> str:
    candidates: list[Path] = []
    found = shutil.which("pdftoppm")
    if found:
        wrapper = Path(found)
        candidates.append(wrapper)
        for parent in wrapper.parents:
            candidates.append(parent / "native" / "poppler" / "Library" / "bin" / "pdftoppm.exe")
            candidates.append(parent / "native" / "poppler" / "bin" / "pdftoppm.exe")
    for candidate in candidates:
        if candidate.is_file() and candidate.suffix.lower() == ".exe":
            return str(candidate)
    return str(candidates[0]) if candidates else ""


class RenderedPages(list[Path]):
    def __init__(self, pages: list[Path] | None = None, error: str = "") -> None:
        super().__init__(pages or [])
        self.error = error


def render_pdf_for_ocr(path: Path, output_dir: Path) -> RenderedPages:
    pdftoppm = locate_pdftoppm()
    if not pdftoppm:
        return RenderedPages(error="pdftoppm executable not found; install Poppler or add pdftoppm to PATH")
    prefix = output_dir / "ocr_page"
    try:
        completed = subprocess.run(
            [pdftoppm, "-png", "-r", "180", str(path), str(prefix)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=600,
        )
    except Exception as exc:
        return RenderedPages(error=f"pdftoppm execution failed: {exc}")
    if completed.returncode:
        detail = completed.stderr.strip() or completed.stdout.strip() or f"exit code {completed.returncode}"
        return RenderedPages(error=f"pdftoppm failed: {detail}")
    pages = sorted(output_dir.glob("ocr_page-*.png"), key=lambda item: int(item.stem.rsplit("-", 1)[-1]))
    return RenderedPages(pages, error="" if pages else "pdftoppm completed without producing PNG pages")


def ocr_image_text(image_path: Path) -> str:
    script = Path(__file__).with_name("ocr_windows.ps1")
    powershell = resolve_powershell_host(os.environ.get(WINRT_OCR_POWERSHELL_HOST_ENV))
    if not script.is_file() or not powershell.path.is_file():
        raise RuntimeError("Windows OCR script or PowerShell executable is unavailable")
    completed = subprocess.run(
        [str(powershell.path), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), "-ImagePath", str(image_path)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
    )
    if completed.returncode != 0:
        detail = completed.stderr.strip() or completed.stdout.strip() or f"exit code {completed.returncode}"
        raise RuntimeError(f"Windows OCR failed: {detail}")
    return completed.stdout


def locate_tesseract() -> str:
    for variable in ("TESSERACT_CMD", "TESSERACT_EXE", "TESSERACT_PATH"):
        configured = str(os.environ.get(variable) or "").strip().strip('"')
        if configured and Path(configured).is_file():
            return configured
    return shutil.which("tesseract") or ""


def ocr_image_text_tesseract(image_path: Path) -> str:
    executable = locate_tesseract()
    if not executable:
        raise RuntimeError("Tesseract executable not found; set TESSERACT_CMD or add tesseract to PATH")
    language = str(os.environ.get("TESSERACT_LANG") or "chi_sim+eng").strip() or "chi_sim+eng"
    completed = subprocess.run(
        [executable, str(image_path), "stdout", "-l", language],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=180,
    )
    if completed.returncode != 0:
        detail = completed.stderr.strip() or completed.stdout.strip() or f"exit code {completed.returncode}"
        raise RuntimeError(f"Tesseract OCR failed: {detail}")
    return completed.stdout


def ocr_image_text_with_fallback(image_path: Path) -> tuple[str, str, dict[str, str]]:
    diagnostics: dict[str, str] = {}
    try:
        text = ocr_image_text(image_path)
        if normalize_text(text):
            return text, "windows", diagnostics
        diagnostics["windows"] = "completed without recognized text"
    except Exception as exc:
        diagnostics["windows"] = str(exc)
    try:
        text = ocr_image_text_tesseract(image_path)
        if normalize_text(text):
            return text, "tesseract", diagnostics
        diagnostics["tesseract"] = "completed without recognized text"
    except Exception as exc:
        diagnostics["tesseract"] = str(exc)
    return "", "", diagnostics


def inventory_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    doc = Document(path)
    collector = ItemCollector()
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = normalize_text(paragraph.text)
        if text:
            location = f"paragraph:{index}"
            collector.add(text_item("text", location, text, style=paragraph.style.name if paragraph.style else ""))
    for table_index, table in enumerate(doc.tables, 1):
        rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
        location = f"table:{table_index}"
        collector.add({"id": stable_id("table", location, json.dumps(rows, ensure_ascii=False)), "kind": "table", "location": location, "rows": rows})
    media_count = 0
    relationship_count = 0
    description_count = 0
    header_count = 0
    footer_count = 0
    metadata = {}
    with zipfile.ZipFile(path) as package:
        for name in sorted(package.namelist()):
            if re.fullmatch(r"word/media/[^/]+", name):
                data = package.read(name)
                if collector.add(media_item("image", name, Path(name).name, data)):
                    media_count += 1
            if name.startswith("docProps/") and name.endswith(".xml"):
                metadata[name] = " ".join(etree.fromstring(package.read(name)).itertext())
            if re.fullmatch(r"word/(header|footer)[^/]*\.xml", name):
                region = "header" if "/header" in name else "footer"
                part_root = etree.fromstring(package.read(name))
                tables = [element for element in part_root.iter() if local_name(element) == "tbl"]
                for table_index, table in enumerate(tables, 1):
                    rows = []
                    for row in (element for element in table if local_name(element) == "tr"):
                        cells = []
                        for cell in (element for element in row if local_name(element) == "tc"):
                            cells.append(normalize_text(" ".join(text for text in cell.itertext())))
                        rows.append(cells)
                    location = f"{region}:{name}/table:{table_index}"
                    collector.add({"id": stable_id("table", location, json.dumps(rows, ensure_ascii=False)), "kind": "table", "location": location, "rows": rows, "region": region, "part": name})
                paragraph_index = 0
                for paragraph in (element for element in part_root.iter() if local_name(element) == "p"):
                    if any(local_name(parent) == "tbl" for parent in _ancestors(paragraph)):
                        continue
                    text = normalize_text(" ".join(text for text in paragraph.itertext()))
                    if not text:
                        continue
                    paragraph_index += 1
                    location = f"{region}:{name}/paragraph:{paragraph_index}"
                    if collector.add(text_item("text", location, text, region=region, part=name)):
                        if region == "header":
                            header_count += 1
                        else:
                            footer_count += 1
                if region == "header":
                    header_count += len(tables)
                else:
                    footer_count += len(tables)
            if name.startswith("word/") and name.endswith(".xml") and not name.endswith(".rels"):
                root = etree.fromstring(package.read(name))
                description_count += add_ooxml_descriptions(root, name, collector)
                for textbox_index, textbox in enumerate((element for element in root.iter() if local_name(element) == "txbxContent"), 1):
                    for paragraph_index, paragraph in enumerate((element for element in textbox.iter() if local_name(element) == "p"), 1):
                        text = normalize_text(" ".join(text for text in paragraph.itertext()))
                        if text:
                            location = f"textbox:{name}:{textbox_index}/paragraph:{paragraph_index}"
                            collector.add(text_item("text", location, text, region="textbox", part=name))
        relationship_count = add_external_relationship_items(package, collector)
    return {
        "format": "docx",
        "items": collector.items,
        "metadata": metadata,
        "counts": {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "media": media_count, "headers": header_count, "footers": footer_count, "headers_footers": header_count + footer_count, "relationships": relationship_count, "external_relationships": relationship_count, "descriptions": description_count},
    }


def inventory_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    collector = ItemCollector()
    seen_media_objects: set[tuple[int, str]] = set()
    seen_media_hashes: set[str] = set()

    def visit(shape, location: str, slide_index: int):
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            for child_index, child in enumerate(shape.shapes, 1):
                visit(child, f"{location}/group:{child_index}", slide_index)
            return
        if getattr(shape, "has_text_frame", False):
            text = normalize_text(shape.text)
            if text:
                collector.add(text_item("text", location, text, slide=slide_index))
        if getattr(shape, "has_table", False):
            rows = [[normalize_text(cell.text) for cell in row.cells] for row in shape.table.rows]
            collector.add({"id": stable_id("table", location, json.dumps(rows, ensure_ascii=False)), "kind": "table", "location": location, "rows": rows, "slide": slide_index})
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
            data = shape.image.blob
            seen_media_hashes.add(hashlib.sha256(data).hexdigest().upper())
            embed = ""
            for element in shape._element.iter():
                if local_name(element) == "blip":
                    embed = str(element.get(f"{{{OFFICE_REL_NS}}}embed") or "")
                    break
            media_key = (slide_index, embed) if embed else (slide_index, location)
            if media_key in seen_media_objects:
                return
            seen_media_objects.add(media_key)
            media_location = f"slide:{slide_index}/media:{embed}" if embed else location
            collector.add(media_item("image", media_location, f"slide{slide_index}_{location.rsplit(':', 1)[-1]}.{shape.image.ext}", data, slide=slide_index, relationship_id=embed))
    for slide_index, slide in enumerate(presentation.slides, 1):
        for shape_index, shape in enumerate(slide.shapes, 1):
            location = f"slide:{slide_index}/shape:{shape_index}"
            visit(shape, location, slide_index)

    notes_count = 0
    master_count = 0
    layout_count = 0
    description_count = 0
    relationship_count = 0
    package_media_count = 0
    unreferenced_package_media_count = 0
    with zipfile.ZipFile(path) as package:
        for name in sorted(package.namelist()):
            if re.fullmatch(r"ppt/media/[^/]+", name):
                package_media_count += 1
                data = package.read(name)
                digest = hashlib.sha256(data).hexdigest().upper()
                if digest not in seen_media_hashes:
                    kind = "image" if image_visual_sha256(data) else "media"
                    if collector.add(media_item(kind, f"package:{name}", Path(name).name, data, region="package_media", part=name)):
                        unreferenced_package_media_count += 1
                continue
            if not name.startswith("ppt/") or not name.endswith(".xml"):
                continue
            root = etree.fromstring(package.read(name))
            description_count += add_ooxml_descriptions(root, name, collector)
            region = ""
            if "/notesSlides/" in name:
                region = "notes"
                notes_count += 1
            elif "/slideMasters/" in name:
                region = "slide_master"
                master_count += 1
            elif "/slideLayouts/" in name:
                region = "slide_layout"
                layout_count += 1
            if region:
                paragraph_index = 0
                for paragraph in (element for element in root.iter() if local_name(element) == "p"):
                    text = normalize_text(" ".join(text for text in paragraph.itertext()))
                    if not text:
                        continue
                    paragraph_index += 1
                    location = f"{region}:{name}/paragraph:{paragraph_index}"
                    collector.add(text_item("text", location, text, region=region, part=name))
        relationship_count = add_external_relationship_items(package, collector)
    return {
        "format": "pptx",
        "items": collector.items,
        "metadata": {},
        "counts": {"slides": len(presentation.slides), "items": len(collector.items), "notes": notes_count, "notes_slides": notes_count, "slide_masters": master_count, "slide_layouts": layout_count, "masters": master_count, "layouts": layout_count, "descriptions": description_count, "relationships": relationship_count, "external_relationships": relationship_count, "package_media": package_media_count, "unreferenced_package_media": unreferenced_package_media_count},
    }


def inventory_pdf(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    collector = ItemCollector()
    findings: list[dict[str, Any]] = []
    pages_without_text: list[int] = []
    for page_index, page in enumerate(reader.pages, 1):
        try:
            lines = (page.extract_text() or "").splitlines()
        except Exception as exc:
            lines = []
            findings.append({"code": "PDF_TEXT_EXTRACTION_FAILED", "page": page_index, "message": str(exc)})
        page_has_text = False
        for line_index, line in enumerate(lines, 1):
            text_items: list[dict[str, Any]] = []
            page_has_text = add_pdf_text_item(text_items, page_index, line_index, line) or page_has_text
            for item in text_items:
                collector.add(item)
        if not page_has_text:
            pages_without_text.append(page_index)
        for image_index, image in enumerate(page.images, 1):
            data = image.data
            extension = Path(image.name).suffix or ".bin"
            location = f"page:{page_index}/image:{image_index}"
            collector.add(media_item("image", location, f"page{page_index}_{image_index}{extension}", data, page=page_index))
    ocr_pages: list[int] = []
    ocr_methods: dict[str, str] = {}
    if pages_without_text:
        with tempfile.TemporaryDirectory(prefix="zxty-pdf-ocr-") as directory:
            try:
                rendered_pages = render_pdf_for_ocr(path, Path(directory))
            except Exception as exc:
                rendered_pages = RenderedPages(error=f"PDF OCR rendering raised an exception: {exc}")
            rendered_by_page: dict[int, Path] = {}
            for rendered in rendered_pages:
                try:
                    rendered_by_page[int(rendered.stem.rsplit("-", 1)[-1])] = rendered
                except (AttributeError, ValueError):
                    continue
            for page_index in pages_without_text:
                png = rendered_by_page.get(page_index)
                if png is None:
                    findings.append({
                        "code": "PDF_OCR_RENDER_FAILED",
                        "page": page_index,
                        "message": getattr(rendered_pages, "error", "required page was not rendered"),
                    })
                    continue
                text, method, diagnostics = ocr_image_text_with_fallback(png)
                if text and add_pdf_ocr_text_item(collector.items, page_index, text):
                    item = collector.items[-1]
                    item["source_id"] = item["id"]
                    ocr_pages.append(page_index)
                    ocr_methods[str(page_index)] = method
                    continue
                findings.append({
                    "code": "PDF_OCR_FAILED",
                    "page": page_index,
                    "image": str(png),
                    "attempts": diagnostics,
                })
    metadata = {str(key): str(value) for key, value in (reader.metadata or {}).items()}
    return {
        "status": "BLOCKED" if findings else "PASS",
        "findings": findings,
        "format": "pdf",
        "items": collector.items,
        "metadata": metadata,
        "counts": {"pages": len(reader.pages), "items": len(collector.items), "ocr_pages": len(ocr_pages), "ocr_required_pages": len(pages_without_text), "ocr_failed_pages": len(pages_without_text) - len(ocr_pages)},
        "ocr_pages": ocr_pages,
        "ocr_methods": ocr_methods,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="生成源资料完整内容清单")
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    source = args.source.resolve()
    if not source.is_file():
        payload = result("BLOCKED", "inventory_source", findings=[{"code": "SOURCE_MISSING"}], source=str(source))
        write_json(args.output, payload)
        return finish(payload, args.report)
    try:
        handlers = {".docx": inventory_docx, ".pptx": inventory_pptx, ".pdf": inventory_pdf}
        handler = handlers.get(source.suffix.lower())
        if not handler:
            payload = result("BLOCKED", "inventory_source", findings=[{"code": "NORMALIZATION_REQUIRED", "message": source.suffix}], source=str(source))
            write_json(args.output, payload)
            return finish(payload, args.report)
        facts = handler(source)
        status = str(facts.pop("status", "PASS"))
        findings = list(facts.pop("findings", []))
        payload = result(status, "inventory_source", findings=findings, source=str(source), source_sha256=sha256_file(source), **facts)
    except Exception as exc:
        payload = result("BLOCKED", "inventory_source", findings=[{"code": "INVENTORY_FAILED", "message": str(exc)}], source=str(source))
    write_json(args.output, payload)
    return finish(payload, args.report)


if __name__ == "__main__":
    raise SystemExit(main())
