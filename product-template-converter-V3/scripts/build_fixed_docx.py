from __future__ import annotations

import argparse
import copy
import io
import posixpath
import re
from collections import deque
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips
from PIL import Image
from lxml import etree

from pipeline_common import SECTIONS, finish, read_json, result, sha256_file


TEMPLATE_SHA256 = "90B04F4C802A3821234B7A47789CCB47904C7D8E5C486087EC54953A39D19E70"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
EMU_PER_INCH = 914400
PIXELS_PER_INCH = 96


class ImageMaterial:
    def __init__(self, data: bytes, source_display_size: tuple[float, float] | None = None):
        self.data = data
        self.source_display_size = source_display_size


def set_font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def clear_para(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)




def add_cover_body_page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


def verify_cover_body_page_break(path: Path):
    with zipfile.ZipFile(path) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    body_node = root.find("w:body", ns)
    body_start = None
    page_break_blocks = []
    for index, block in enumerate(list(body_node)):
        text = "".join(block.xpath(".//w:t/text()", namespaces=ns)).strip()
        if block.xpath('.//w:br[@w:type="page"]', namespaces=ns):
            page_break_blocks.append(index)
        if text.startswith("▌ 一、产品简介") or text.startswith("一、产品简介"):
            body_start = index
            break
    if body_start is None:
        raise ValueError("未找到正文起始章节：一、产品简介")
    if not page_break_blocks or page_break_blocks[-1] >= body_start:
        raise ValueError("封面与正文之间缺少显式分页符")


def set_outline(paragraph, level):
    ppr = paragraph._p.get_or_add_pPr()
    for tag in ("keepNext", "keepLines", "pageBreakBefore", "numPr", "pBdr", "outlineLvl"):
        node = ppr.find(qn(f"w:{tag}"))
        if node is not None:
            ppr.remove(node)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    ppr.append(outline)


def heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(6 if level == 1 else 3)
    set_outline(paragraph, level - 1)
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "DCE6F1")
        paragraph._p.get_or_add_pPr().append(shade)
        set_font(paragraph.add_run("▌ " + text), 14, True, "1F4E79")
    else:
        set_font(paragraph.add_run("▎" + text), 11, True, "1F4E79")
    return paragraph


def body(doc, text, compact=False, keep_with_next=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.space_after = Pt(2 if compact else 5)
    paragraph.paragraph_format.line_spacing = 1.1 if compact else 1.3
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = True
    set_font(paragraph.add_run(text), 9.5 if compact else 10.5)
    return paragraph


def term(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    set_font(paragraph.add_run(text), 10.5)


def apply_redactions(text: str, entry: dict) -> str:
    for value in entry.get("redactions", []):
        text = text.replace(str(value), "")
    return re.sub(r"\s+", " ", text).strip()


def text_for_entry(item: dict, entry: dict) -> str:
    if source_title_consumed_by_template(item, entry):
        return ""
    if entry.get("action") == "reviewed_text":
        return re.sub(r"\s+", " ", str(entry.get("reviewed_text", ""))).strip()
    return apply_redactions(item.get("text", ""), entry) if entry.get("action") == "redact_identity" else item.get("text", "")


def _section_title_token(text: str) -> str:
    return re.sub(r"[\s▌▎]+", "", str(text or ""))


def source_title_consumed_by_template(item: dict, entry: dict) -> bool:
    """Return true when a source heading is represented by the fixed template."""

    if item.get("kind") != "text" or entry.get("target_section") not in SECTIONS:
        return False
    style = str(item.get("style") or "").casefold()
    if "heading" not in style:
        return False
    source_text = _section_title_token(item.get("text", ""))
    if not source_text:
        return False
    return (
        source_text in {_section_title_token(section) for section in SECTIONS}
        or bool(re.match(r"^[一二三四五六七八九十]+[、.．]", source_text))
    )


def replacement_path(entry: dict, content_map_path: Path) -> Path | None:
    raw_path = entry.get("replacement_path")
    if not isinstance(raw_path, str) or not raw_path.strip():
        return None
    path = Path(raw_path)
    return path if path.is_absolute() else content_map_path.parent / path


def sanitized_image_findings(entry: dict, item: dict, content_map_path: Path) -> list[dict]:
    source_id = entry.get("source_id", "")
    findings = []
    if item.get("kind") != "image":
        findings.append({"code": "PRESERVE_SANITIZED_IMAGE_REQUIRES_IMAGE", "source_id": source_id})
    replacement = replacement_path(entry, content_map_path)
    if replacement is None:
        findings.append({"code": "SANITIZED_IMAGE_REPLACEMENT_PATH_MISSING", "source_id": source_id})
    elif not replacement.is_file():
        findings.append({"code": "SANITIZED_IMAGE_REPLACEMENT_PATH_NOT_FOUND", "source_id": source_id, "path": str(replacement)})
    expected_sha = entry.get("replacement_sha256")
    if not isinstance(expected_sha, str) or not expected_sha.strip():
        findings.append({"code": "SANITIZED_IMAGE_REPLACEMENT_SHA256_MISSING", "source_id": source_id})
    elif replacement is not None and replacement.is_file():
        actual_sha = sha256_file(replacement)
        if actual_sha != expected_sha.strip().upper():
            findings.append({"code": "SANITIZED_IMAGE_REPLACEMENT_SHA256_MISMATCH", "source_id": source_id, "expected": expected_sha.strip().upper(), "actual": actual_sha})
    if not isinstance(entry.get("sanitization_review"), str) or not entry.get("sanitization_review", "").strip():
        findings.append({"code": "SANITIZED_IMAGE_REVIEW_MISSING", "source_id": source_id})
    if entry.get("target_section") not in SECTIONS:
        findings.append({"code": "INVALID_TARGET_SECTION", "source_id": source_id, "target_section": entry.get("target_section", "")})
    return findings


def apply_sanitized_image_replacements(images: dict[str, ImageMaterial], entries: list[dict], content_map_path: Path) -> None:
    """把已校验的替换图绑定到源 ID，后续封面和正文都只能读取这份字节。"""
    for entry in entries:
        if entry.get("action") != "preserve_sanitized_image":
            continue
        replacement = replacement_path(entry, content_map_path)
        if replacement is None:
            raise ValueError(f"缺少脱敏图片替换路径: {entry.get('source_id', '')}")
        existing = images.get(entry["source_id"])
        display_size = existing.source_display_size if isinstance(existing, ImageMaterial) else None
        images[entry["source_id"]] = ImageMaterial(replacement.read_bytes(), display_size)


def expand_reviewed_text_blocks(entries: list[dict]) -> list[dict]:
    expanded = []
    for entry in entries:
        blocks = entry.get("reviewed_blocks")
        if entry.get("action") != "reviewed_text" or not isinstance(blocks, list):
            expanded.append(entry)
            continue
        for index, block in enumerate(blocks, 1):
            expanded.append({
                **entry,
                "reviewed_blocks": None,
                "reviewed_text": block.get("text", ""),
                "target_section": block.get("target_section", ""),
                "target_order": block.get("target_order", 0),
                "source_block_index": index,
            })
    return expanded


_SERIAL_COLUMN_HEADERS = {"序号", "序列号", "编号", "no"}
_SERIAL_VALUE_RE = re.compile(r"^(?:\d+|[一二三四五六七八九十百千零〇]+|[①②③④⑤⑥⑦⑧⑨⑩]+|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+)$")


def _normalize_serial_token(text: str) -> str:
    return re.sub(r"[\s\.\．、,，:：;；\-_/\\\(\)（）\[\]【】]+", "", str(text)).lower()


def _is_serial_column_header(text: str) -> bool:
    return _normalize_serial_token(text) in _SERIAL_COLUMN_HEADERS


def _is_serial_column_value(text: str) -> bool:
    token = _normalize_serial_token(text)
    return bool(token) and bool(_SERIAL_VALUE_RE.fullmatch(token))


def drop_serial_parameter_column(rows):
    normalized_rows = [list(row) if isinstance(row, (list, tuple)) else [row] for row in rows]
    if len(normalized_rows) < 2:
        return normalized_rows

    first_column = [str(row[0]).strip() if row else "" for row in normalized_rows]
    first_nonempty = next((value for value in first_column if value), "")
    if not first_nonempty:
        return normalized_rows

    rest_has_content = any(len(row) > 1 and any(str(cell).strip() for cell in row[1:]) for row in normalized_rows)
    if _is_serial_column_header(first_nonempty):
        serial_values = [value for value in first_column[1:] if value]
        if serial_values and all(_is_serial_column_value(value) for value in serial_values):
            return [list(row[1:]) for row in normalized_rows]
        return normalized_rows

    if rest_has_content and all(not value or _is_serial_column_value(value) for value in first_column):
        return [list(row[1:]) for row in normalized_rows]

    return normalized_rows


def qualification_slot_lines(section_entries: list[dict]) -> list[str]:
    """六、产品资质不合成缺省文案。"""
    if not section_entries:
        return []
    return []


def should_render_section(section: str, section_entries: list[dict]) -> bool:
    """固定模板始终保留六篇章标题；空章节不合成来源之外的正文。"""
    return section in SECTIONS


def use_compact_body(section: str, section_entries: list[dict]) -> bool:
    """密集参数列表减少段间空白，同时保留每条参数不跨页。"""
    return section == "五、技术参数" and len(section_entries) >= 10


def prevent_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:tblHeader")) is None:
        node = OxmlElement("w:tblHeader")
        node.set(qn("w:val"), "true")
        trpr.append(node)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, column_count):
    """显式锁定表格网格，避免 Word/WPS/手机查看器各自自动分列。"""
    total_width = 8300
    if column_count == 3:
        widths = [900, 1800, total_width - 2700]
    elif column_count == 2:
        widths = [2200, total_width - 2200]
    else:
        base, remainder = divmod(total_width, column_count)
        widths = [base + (1 if index < remainder else 0) for index in range(column_count)]
    tbl_pr = table._tbl.tblPr
    table_width = tbl_pr.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(total_width))
    table_width.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def trim_white_background_image(image: Image.Image, background_threshold: int = 248, margin_px: int = 4) -> tuple[Image.Image, bool]:
    rgba = image.convert("RGBA")
    width, height = rgba.size
    if width == 0 or height == 0:
        return rgba, False

    pixels = rgba.load()
    visited = bytearray(width * height)
    queue = deque()

    def is_background(x: int, y: int) -> bool:
        r, g, b, a = pixels[x, y]
        return a == 0 or (r >= background_threshold and g >= background_threshold and b >= background_threshold)

    def enqueue(x: int, y: int) -> None:
        idx = y * width + x
        if not visited[idx] and is_background(x, y):
            queue.append((x, y))

    for x in range(width):
        enqueue(x, 0)
        enqueue(x, height - 1)
    for y in range(height):
        enqueue(0, y)
        enqueue(width - 1, y)

    while queue:
        x, y = queue.popleft()
        idx = y * width + x
        if visited[idx]:
            continue
        visited[idx] = 1
        if not is_background(x, y):
            continue
        if x > 0:
            enqueue(x - 1, y)
        if x + 1 < width:
            enqueue(x + 1, y)
        if y > 0:
            enqueue(x, y - 1)
        if y + 1 < height:
            enqueue(x, y + 1)

    left, top = width, height
    right, bottom = -1, -1
    for y in range(height):
        row_offset = y * width
        for x in range(width):
            if visited[row_offset + x]:
                continue
            if x < left:
                left = x
            if y < top:
                top = y
            if x > right:
                right = x
            if y > bottom:
                bottom = y

    if right < left or bottom < top:
        return rgba, False

    left = max(0, left - margin_px)
    top = max(0, top - margin_px)
    right = min(width - 1, right + margin_px)
    bottom = min(height - 1, bottom + margin_px)
    if left == 0 and top == 0 and right == width - 1 and bottom == height - 1:
        return rgba, False
    return rgba.crop((left, top, right + 1, bottom + 1)), True


def add_table(doc, rows):
    width = max((len(row) for row in rows), default=2)
    if width < 2:
        width = 2
    table = doc.add_table(rows=0, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_split(table.rows[-1])
        if row_index == 0:
            repeat_header(table.rows[-1])
        for index, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = values[index] if index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            # 小字号表格不吸附正文行网格；仅表头跟随下一行，避免孤立表头。
            snap_to_grid = OxmlElement("w:snapToGrid")
            snap_to_grid.set(qn("w:val"), "0")
            paragraph._p.get_or_add_pPr().append(snap_to_grid)
            paragraph.paragraph_format.keep_with_next = row_index == 0
            set_font(paragraph.add_run(text), 9.5, row_index == 0, "FFFFFF" if row_index == 0 else "000000")
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "1F4E79")
                cell._tc.get_or_add_tcPr().append(shade)
    set_table_geometry(table, width)
    for row, values in zip(table.rows, rows):
        nonempty = [index for index, value in enumerate(values) if str(value).strip()]
        if nonempty and nonempty[-1] < width - 1:
            row.cells[nonempty[-1]].merge(row.cells[-1])
    return table


def _valid_display_size(width: float | None, height: float | None) -> tuple[float, float] | None:
    if width is None or height is None or width <= 0 or height <= 0:
        return None
    return float(width), float(height)


def _smaller_display_size(previous: tuple[float, float] | None, candidate: tuple[float, float]) -> tuple[float, float]:
    if previous is None or candidate[0] * candidate[1] < previous[0] * previous[1]:
        return candidate
    return previous


def _css_length_to_inches(value: str | None) -> float | None:
    if not value:
        return None
    match = re.fullmatch(r"\s*([0-9]+(?:\.[0-9]+)?)\s*([a-zA-Z]*)\s*", value)
    if not match:
        return None
    amount = float(match.group(1))
    unit = match.group(2).lower() or "pt"
    if unit == "in":
        return amount
    if unit == "pt":
        return amount / 72
    if unit == "cm":
        return amount / 2.54
    if unit == "mm":
        return amount / 25.4
    if unit == "px":
        return amount / PIXELS_PER_INCH
    return None


def _vml_shape_display_size(shape) -> tuple[float, float] | None:
    style = str(shape.get("style") or "")
    declarations = {}
    for declaration in style.split(";"):
        if ":" not in declaration:
            continue
        name, value = declaration.split(":", 1)
        declarations[name.strip().lower()] = value.strip()
    return _valid_display_size(
        _css_length_to_inches(declarations.get("width")),
        _css_length_to_inches(declarations.get("height")),
    )


def docx_image_display_sizes(package: zipfile.ZipFile) -> dict[str, tuple[float, float]]:
    relationship_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
    word_ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "o": "urn:schemas-microsoft-com:office:office",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "v": "urn:schemas-microsoft-com:vml",
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": WP_NS,
    }
    package_names = set(package.namelist())
    by_media: dict[str, tuple[float, float]] = {}

    def relationships_for_part(part_name: str) -> dict[str, str]:
        directory = posixpath.dirname(part_name)
        rels_name = posixpath.join(directory, "_rels", posixpath.basename(part_name) + ".rels")
        if rels_name not in package_names:
            return {}
        root = etree.fromstring(package.read(rels_name))
        return {
            str(node.get("Id")): posixpath.normpath(posixpath.join(directory, str(node.get("Target") or "")))
            for node in root.xpath("./r:Relationship[not(@TargetMode='External')]", namespaces=relationship_ns)
        }

    for part_name in sorted(name for name in package_names if name.startswith("word/") and name.endswith(".xml") and "/_rels/" not in name):
        relationships = relationships_for_part(part_name)
        if not relationships:
            continue
        try:
            root = etree.fromstring(package.read(part_name))
        except etree.XMLSyntaxError:
            continue
        for drawing in root.xpath(".//w:drawing", namespaces=word_ns):
            blips = drawing.xpath(".//a:blip", namespaces=word_ns)
            extents = drawing.xpath(".//wp:extent", namespaces=word_ns)
            if not blips or not extents:
                continue
            relationship_id = str(blips[0].get(f"{{{word_ns['r']}}}embed") or "")
            media = relationships.get(relationship_id)
            if not media or media not in package_names:
                continue
            width = int(extents[0].get("cx") or 0) / EMU_PER_INCH
            height = int(extents[0].get("cy") or 0) / EMU_PER_INCH
            display_size = _valid_display_size(width, height)
            if display_size is None:
                continue
            by_media[media] = _smaller_display_size(by_media.get(media), display_size)
        for pict in root.xpath(".//w:pict", namespaces=word_ns):
            image_nodes = pict.xpath(".//v:imagedata", namespaces=word_ns)
            shape_nodes = pict.xpath(".//v:shape", namespaces=word_ns)
            if not image_nodes or not shape_nodes:
                continue
            relationship_id = str(
                image_nodes[0].get(f"{{{word_ns['r']}}}id")
                or image_nodes[0].get(f"{{{word_ns['o']}}}relid")
                or ""
            )
            media = relationships.get(relationship_id)
            if not media or media not in package_names:
                continue
            display_size = _vml_shape_display_size(shape_nodes[0])
            if display_size is None:
                continue
            by_media[media] = _smaller_display_size(by_media.get(media), display_size)
    return by_media


def pptx_image_display_sizes(package: zipfile.ZipFile) -> dict[tuple[int, str], tuple[float, float]]:
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    display_sizes: dict[tuple[int, str], tuple[float, float]] = {}

    def group_scale(picture) -> tuple[float, float] | None:
        scale_x = 1.0
        scale_y = 1.0
        for group in picture.xpath("ancestor::p:grpSp", namespaces=namespaces):
            extents = group.xpath("./p:grpSpPr/a:xfrm/a:ext", namespaces=namespaces)
            child_extents = group.xpath("./p:grpSpPr/a:xfrm/a:chExt", namespaces=namespaces)
            if not extents or not child_extents:
                return None
            ext_width = int(extents[0].get("cx") or 0)
            ext_height = int(extents[0].get("cy") or 0)
            child_width = int(child_extents[0].get("cx") or 0)
            child_height = int(child_extents[0].get("cy") or 0)
            if ext_width <= 0 or ext_height <= 0 or child_width <= 0 or child_height <= 0:
                return None
            scale_x *= ext_width / child_width
            scale_y *= ext_height / child_height
        return scale_x, scale_y

    for slide_name in sorted(name for name in package.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)):
        slide_match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", slide_name)
        if not slide_match:
            continue
        slide_index = int(slide_match.group(1))
        try:
            root = etree.fromstring(package.read(slide_name))
        except etree.XMLSyntaxError:
            continue
        for picture in root.xpath(".//p:pic", namespaces=namespaces):
            blips = picture.xpath(".//a:blip", namespaces=namespaces)
            extents = picture.xpath("./p:spPr/a:xfrm/a:ext", namespaces=namespaces)
            if not blips or not extents:
                continue
            relationship_id = str(blips[0].get(f"{{{namespaces['r']}}}embed") or "")
            scale = group_scale(picture)
            if scale is None:
                continue
            width = int(extents[0].get("cx") or 0) * scale[0] / EMU_PER_INCH
            height = int(extents[0].get("cy") or 0) * scale[1] / EMU_PER_INCH
            display_size = _valid_display_size(width, height)
            if not relationship_id or display_size is None:
                continue
            key = (slide_index, relationship_id)
            display_sizes[key] = _smaller_display_size(display_sizes.get(key), display_size)
    return display_sizes


def load_images(inventory: dict, source_path: Path) -> dict[str, ImageMaterial]:
    blobs: dict[str, ImageMaterial] = {}
    suffix = source_path.suffix.lower()
    if suffix == ".docx":
        with zipfile.ZipFile(source_path) as package:
            display_sizes = docx_image_display_sizes(package)
            for item in inventory.get("items", []):
                if item.get("kind") == "image" and item.get("location") in package.namelist():
                    location = item["location"]
                    blobs[item["id"]] = ImageMaterial(package.read(location), display_sizes.get(location))
    elif suffix == ".pptx":
        relationship_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        with zipfile.ZipFile(source_path) as package:
            package_names = set(package.namelist())
            relationship_cache: dict[int, dict[str, str]] = {}
            display_sizes = pptx_image_display_sizes(package)
            for item in inventory.get("items", []):
                if item.get("kind") != "image":
                    continue
                match = re.fullmatch(r"slide:(\d+)/media:(rId[^/]+)", str(item.get("location", "")))
                if not match:
                    continue
                slide_index = int(match.group(1))
                relationship_id = match.group(2)
                if slide_index not in relationship_cache:
                    rels_name = f"ppt/slides/_rels/slide{slide_index}.xml.rels"
                    if rels_name not in package_names:
                        relationship_cache[slide_index] = {}
                    else:
                        root = etree.fromstring(package.read(rels_name))
                        relationship_cache[slide_index] = {
                            str(node.get("Id")): str(node.get("Target"))
                            for node in root.xpath("./r:Relationship[not(@TargetMode='External')]", namespaces=relationship_ns)
                        }
                target = relationship_cache[slide_index].get(relationship_id, "")
                member = posixpath.normpath(posixpath.join("ppt/slides", target)) if target else ""
                if member in package_names:
                    blobs[item["id"]] = ImageMaterial(package.read(member), display_sizes.get((slide_index, relationship_id)))
    elif suffix == ".pdf":
        from pypdf import PdfReader
        from source_image_geometry_v3 import pdf_image_display_sizes

        document = PdfReader(source_path)
        display_sizes = pdf_image_display_sizes(source_path)
        for item in inventory.get("items", []):
            if item.get("kind") != "image":
                continue
            found = re.fullmatch(r"page:(\d+)/image:(\d+)", item.get("location", ""))
            if found:
                page_index, image_index = int(found.group(1)) - 1, int(found.group(2)) - 1
                blobs[item["id"]] = ImageMaterial(document.pages[page_index].images[image_index].data, display_sizes.get((page_index + 1, image_index + 1)))
    return blobs


def source_image_display_sizes(inventory: dict, source_path: Path) -> dict[str, tuple[float, float]]:
    return {
        source_id: material.source_display_size
        for source_id, material in load_images(inventory, source_path).items()
        if material.source_display_size is not None
    }


def word_compatible_image(data: bytes) -> tuple[bytes, str]:
    """把 JP2/CMYK 等 Word/WPS 不稳定图片统一转成 PNG，保留原始画布。"""
    with Image.open(io.BytesIO(data)) as image:
        original_format = (image.format or "").upper()
        if original_format in {"PNG", "JPEG", "JPG", "GIF", "BMP", "TIFF"} and image.mode not in {"CMYK", "LAB", "P"}:
            return data, original_format
        if image.mode not in {"RGB", "RGBA"}:
            image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
        output = io.BytesIO()
        image.save(output, format="PNG", optimize=True)
        return output.getvalue(), "PNG"


def image_display_size(image: Image.Image, max_width: float, max_height: float, source_display_size: tuple[float, float] | None = None) -> tuple[float, float]:
    ratio = image.width / max(image.height, 1)
    natural_width = image.width / PIXELS_PER_INCH
    natural_height = image.height / PIXELS_PER_INCH
    preferred = _valid_display_size(*(source_display_size or (None, None))) or (natural_width, natural_height)
    width, height = preferred
    if width <= 0 or height <= 0:
        width, height = natural_width, natural_height
    scale = min(1.0, max_width / width if width else 1.0, max_height / height if height else 1.0)
    width *= scale
    height *= scale
    if width <= 0 or height <= 0:
        width = min(max_width, natural_width)
        height = width / ratio
    return width, height


def add_image_to_run(run, material: ImageMaterial | bytes, max_width: float, max_height: float) -> bool:
    try:
        if isinstance(material, ImageMaterial):
            data = material.data
            source_display_size = material.source_display_size
        else:
            data = material
            source_display_size = None
        compatible, _ = word_compatible_image(data)
        with Image.open(io.BytesIO(compatible)) as image:
            width, height = image_display_size(image, max_width, max_height, source_display_size)
        run.add_picture(io.BytesIO(compatible), width=Inches(width), height=Inches(height))
        return True
    except Exception:
        return False


def image_paragraph(doc, data: bytes, max_width=5.7):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    return add_image_to_run(paragraph.add_run(), data, max_width, 4.8)


def image_grid(doc, entries: list[dict], images: dict[str, bytes]) -> list[str]:
    failures = []
    if len(entries) == 1:
        source_id = entries[0]["source_id"]
        if not image_paragraph(doc, images.get(source_id, b"")):
            failures.append(source_id)
        return failures
    for start in range(0, len(entries), 2):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        for index, entry in enumerate(entries[start:start + 2]):
            source_id = entry["source_id"]
            if index:
                paragraph.add_run("   ")
            if not add_image_to_run(paragraph.add_run(), images.get(source_id, b""), 2.65, 3.2):
                failures.append(source_id)
    return failures


def template_closing_bytes(template: Path) -> bytes:
    """固定模板的 image4.jpeg 是经模板哈希锁定的整页封底。"""
    with zipfile.ZipFile(template) as package:
        name = "word/media/image4.jpeg"
        if name not in package.namelist():
            raise ValueError(f"模板缺少固定封底媒体: {name}")
        return package.read(name)


def template_cover_bytes(template: Path) -> bytes:
    """固定模板的 image2.jpeg 是经模板哈希锁定的整页封面背景。"""
    with zipfile.ZipFile(template) as package:
        name = "word/media/image2.jpeg"
        if name not in package.namelist():
            raise ValueError(f"模板缺少固定封面媒体: {name}")
        return package.read(name)


def make_full_page_background(run, section, image_data: bytes):
    run.add_picture(io.BytesIO(image_data), width=section.page_width, height=section.page_height)
    drawing = run._r.xpath(".//w:drawing")[0]
    inline = drawing.xpath("./wp:inline")[0]
    anchor = OxmlElement("wp:anchor")
    for name, value in (
        ("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
        ("simplePos", "0"), ("relativeHeight", "251660288"), ("behindDoc", "1"),
        ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1"),
    ):
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis, align in (("H", "left"), ("V", "top")):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        aligned = OxmlElement("wp:align")
        aligned.text = align
        position.append(aligned)
        anchor.append(position)
    anchor.append(copy.deepcopy(inline.find(qn("wp:extent"))))
    effect = inline.find(qn("wp:effectExtent"))
    if effect is not None:
        anchor.append(copy.deepcopy(effect))
    anchor.append(OxmlElement("wp:wrapNone"))
    for name in ("docPr", "cNvGraphicFramePr"):
        child = inline.find(qn(f"wp:{name}"))
        if child is not None:
            anchor.append(copy.deepcopy(child))
    anchor.append(copy.deepcopy(inline.find(qn("a:graphic"))))
    drawing.remove(inline)
    drawing.append(anchor)


def add_cover_page_background(section, image_data: bytes):
    """封面背景置于首页页眉，避免手机/WPS 重排正文浮动对象。"""
    section.different_first_page_header_footer = True
    header = section.first_page_header
    header.is_linked_to_previous = False
    for paragraph in header.paragraphs:
        clear_para(paragraph)
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(1)
    make_full_page_background(paragraph.add_run(), section, image_data)


def add_closing_page_background(section, image_data: bytes):
    """封底页只能承载固定封底图，正文区保持空载。"""
    section.different_first_page_header_footer = True
    for container in (section.header, section.footer, section.first_page_header, section.first_page_footer):
        container.is_linked_to_previous = False
        for paragraph in container.paragraphs:
            clear_para(paragraph)
    paragraph = section.first_page_header.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(1)
    make_full_page_background(paragraph.add_run(), section, image_data)


def closing_body_placeholder(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(1)
    return paragraph


def following_entry_is_body_image(section_entries: list[dict], index: int, items: dict[str, dict]) -> bool:
    for candidate in section_entries[index + 1:]:
        item = items.get(candidate.get("source_id", ""), {})
        if item.get("kind") == "image":
            if candidate.get("use_on_cover_only"):
                continue
            return True
        return False
    return False


def sanitize_docprops(path: Path):
    """清除来源作者、时间、应用统计和 WPS 自定义记录，不改正文关系。"""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=path.parent) as temporary:
        temp_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    root = etree.fromstring(data)
                    for node in list(root):
                        root.remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                elif info.filename == "docProps/custom.xml":
                    root = etree.fromstring(data)
                    for node in list(root):
                        root.remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                elif info.filename == "docProps/app.xml":
                    root = etree.fromstring(data)
                    for node in list(root):
                        root.remove(node)
                    application = etree.SubElement(root, "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Application")
                    application.text = ""
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                target.writestr(info, data)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def normalize_drawing_ids(path: Path):
    """Assign globally unique wp:docPr ids across body, headers, and footers."""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=path.parent) as temporary:
        temp_path = Path(temporary.name)
    next_id = 1
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                    root = etree.fromstring(data)
                    drawing_properties = root.xpath(".//wp:docPr", namespaces={"wp": WP_NS})
                    if drawing_properties:
                        for node in drawing_properties:
                            node.set("id", str(next_id))
                            next_id += 1
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                target.writestr(info, data)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def create_document(template: Path, source: Path, inventory: dict, mapping: dict, output: Path, content_map_path: Path | None = None) -> list[dict]:
    findings = []
    content_map_path = content_map_path or source.parent / "content_map.json"
    items = {item["id"]: item for item in inventory.get("items", [])}
    for entry in mapping.get("mappings", []):
        if entry.get("action") == "preserve_sanitized_image":
            sanitized_findings = sanitized_image_findings(entry, items.get(entry.get("source_id", ""), {}), content_map_path)
            if sanitized_findings:
                raise ValueError("; ".join(finding["code"] for finding in sanitized_findings))
    doc = Document(template)
    body_root = doc._element.body
    template_sections = body_root.xpath("./w:p/w:pPr/w:sectPr")
    if not template_sections:
        raise ValueError("模板缺少正文节属性")
    first_sect = copy.deepcopy(template_sections[0])
    for child in list(body_root):
        body_root.remove(child)
    body_root.append(OxmlElement("w:p"))
    body_root.append(first_sect)
    doc.sections[0].different_first_page_header_footer = True
    add_cover_page_background(doc.sections[0], template_cover_bytes(template))

    product = mapping["product"]
    doc.add_paragraph().paragraph_format.space_after = Pt(24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"{product['model']} {product['full_name']}"), 19, True, "FFFFFF")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("产 品 介 绍"), 20, True, "FFFFFF")

    images = load_images(inventory, source)
    entries = expand_reviewed_text_blocks(mapping.get("mappings", []))
    apply_sanitized_image_replacements(images, entries, content_map_path)
    cover_entry = next((entry for entry in entries if entry.get("action") in {"preserve_image", "preserve_sanitized_image"} and entry.get("use_on_cover")), None)
    if cover_entry and cover_entry.get("source_id") in images:
        image_paragraph(doc, images[cover_entry["source_id"]], 4.8)
    add_cover_body_page_break(doc)

    by_section: dict[str, list[dict]] = {section: [] for section in SECTIONS}
    for entry in entries:
        if entry.get("action") in {"preserve_exact", "preserve_image", "preserve_sanitized_image", "redact_identity", "reviewed_text"} and entry.get("target_section") in by_section:
            by_section[entry["target_section"]].append(entry)
    for section in SECTIONS:
        section_entries = sorted(by_section[section], key=lambda value: (value.get("target_order", 0), value.get("source_id", ""), value.get("source_block_index", 0)))
        if not should_render_section(section, section_entries):
            continue
        heading(doc, section)
        compact_body = use_compact_body(section, section_entries)
        pending_images = []

        def flush_images():
            for source_id in image_grid(doc, pending_images, images):
                findings.append({"code": "IMAGE_INSERT_FAILED", "source_id": source_id})
            pending_images.clear()

        for entry_index, entry in enumerate(section_entries):
            item = items.get(entry["source_id"], {})
            if item.get("kind") == "image":
                if not entry.get("use_on_cover_only"):
                    pending_images.append(entry)
                continue
            if pending_images:
                flush_images()
            if item.get("kind") == "text":
                text = text_for_entry(item, entry)
                if text:
                    body(
                        doc,
                        text,
                        compact=compact_body,
                        keep_with_next=following_entry_is_body_image(section_entries, entry_index, items),
                    )
            elif item.get("kind") == "table":
                if section == "五、技术参数":
                    rows = drop_serial_parameter_column(item.get("rows", []))
                else:
                    rows = item.get("rows", [])
                add_table(doc, rows)
        if pending_images:
            flush_images()
        if section == "六、产品资质":
            for line in qualification_slot_lines(section_entries):
                body(doc, line)

    closing = doc.add_section(WD_SECTION.NEW_PAGE)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin", "header_distance", "footer_distance"):
        setattr(closing, attr, 0)
    add_closing_page_background(closing, template_closing_bytes(template))
    closing_body_placeholder(doc)

    properties = doc.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""
    properties.title = f"{product['model']} {product['full_name']} 产品介绍"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    normalize_drawing_ids(output)
    sanitize_docprops(output)
    verify_cover_body_page_break(output)
    return findings


def main() -> int:
    parser = argparse.ArgumentParser(description="根据固定模板和已验证映射确定性生成DOCX")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--inventory", type=Path, required=True)
    parser.add_argument("--content-map", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    if sha256_file(args.template) != TEMPLATE_SHA256:
        return finish(result("BLOCKED", "build_fixed_docx", findings=[{"code": "TEMPLATE_HASH_MISMATCH"}]), args.report)
    inventory = read_json(args.inventory)
    mapping = read_json(args.content_map)
    if not mapping.get("product", {}).get("model") or not mapping.get("product", {}).get("full_name"):
        return finish(result("HUMAN_REVIEW", "build_fixed_docx", findings=[{"code": "PRODUCT_IDENTITY_MISSING"}]), args.report)
    sanitized_findings = [
        finding
        for entry in mapping.get("mappings", [])
        if entry.get("action") == "preserve_sanitized_image"
        for finding in sanitized_image_findings(entry, {item["id"]: item for item in inventory.get("items", [])}.get(entry.get("source_id", ""), {}), args.content_map)
    ]
    if sanitized_findings:
        return finish(result("FAIL", "build_fixed_docx", findings=sanitized_findings), args.report)
    try:
        findings = create_document(args.template, args.source, inventory, mapping, args.output, args.content_map)
        status = "HUMAN_REVIEW" if findings else "PASS"
        return finish(result(status, "build_fixed_docx", findings=findings, output=str(args.output), output_sha256=sha256_file(args.output)), args.report)
    except Exception as exc:
        return finish(result("FAIL", "build_fixed_docx", findings=[{"code": "BUILD_FAILED", "message": str(exc)}]), args.report)


if __name__ == "__main__":
    raise SystemExit(main())
