from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image

from pipeline_common import sha256_file
from image_fingerprint import canonical_visual_sha256
from template_ir_v3 import validate_docx_template, validate_template_program


ADAPTER_ID = "declarative-docx-v3"
NON_RENDER_ACTIONS = {"remove_identity", "remove_template_background", "exclude_from_output"}
# 用户 2026-09-05 批准的局部消融范围：仅此模板的空“产品资料补充”。
EMPTY_SUPPLEMENT_TEMPLATE_SHA = "989C35D0EE9C1488CC811B7A0622ACA7764C7274C0866E34CBA1028FB02863E2"
# 新版只修正色带宽度和参数硬分页，保留相同的可选补充槽位；审批仍由模板目录门禁负责。
REPAIRED_SECOND_TEMPLATE_SHA = "54E3CACE204DD1958A2C62991B13C59FB5B22151CC2911AEC674F1811B89224C"


def supplement_has_content(content_plan: dict[str, Any]) -> bool:
    slots = {"body.supplement", "body.supplement.images"}
    for item in content_plan.get("items", []):
        decision = item.get("decision", {})
        if decision.get("action") in NON_RENDER_ACTIONS:
            continue
        blocks = decision.get("reviewed_blocks")
        if isinstance(blocks, list):
            if any(block.get("target_slot") in slots and str(block.get("text") or "").strip() for block in blocks):
                return True
        elif decision.get("target_slot") in slots:
            payload = item.get("source", {}).get("payload", {})
            if payload.get("type") in {"asset", "table"} or str(_text_from_item(item) or "").strip():
                return True
    return False


class DeclarativeAdapterError(ValueError):
    pass


def _slot_locations(slot_contract: dict[str, Any]) -> dict[str, str]:
    records: dict[str, str] = {}
    for slot in slot_contract.get("slots", []):
        if not isinstance(slot, dict) or not isinstance(slot.get("slot_id"), str) or not isinstance(slot.get("location"), str):
            raise DeclarativeAdapterError("SlotContractV3 含无效槽位")
        if slot["slot_id"] in records:
            raise DeclarativeAdapterError(f"SlotContractV3 槽位重复: {slot['slot_id']}")
        records[slot["slot_id"]] = slot["location"]
    return records


def _paragraph_for_slot(document: Document, slot_id: str, locations: dict[str, str]) -> Paragraph:
    location = locations.get(slot_id)
    prefix = "word/document.xml#paragraph:"
    if not location or not location.startswith(prefix):
        raise DeclarativeAdapterError(f"通用 adapter 只接受 document.xml 段落槽位: {slot_id}: {location}")
    try:
        index = int(location[len(prefix):])
    except ValueError as exc:
        raise DeclarativeAdapterError(f"槽位段落索引无效: {slot_id}: {location}") from exc
    paragraphs = document._element.xpath(".//w:p")
    if index < 1 or index > len(paragraphs):
        raise DeclarativeAdapterError(f"槽位段落越界: {slot_id}: {index}/{len(paragraphs)}")
    return Paragraph(paragraphs[index - 1], document._body)


def _replace_paragraph_text(paragraph: Paragraph, lines: list[str]) -> None:
    element = paragraph._p
    # 只替换槽位内容，沿用模板首个文字 run 的直接格式。
    run_properties = element.find('w:r/w:rPr', namespaces=element.nsmap)
    preserved_properties = deepcopy(run_properties) if run_properties is not None else None
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    run = OxmlElement("w:r")
    if preserved_properties is not None:
        run.append(preserved_properties)
    for index, line in enumerate(lines):
        if index:
            run.append(OxmlElement("w:br"))
        text = OxmlElement("w:t")
        if line.startswith(" ") or line.endswith(" "):
            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text.text = line
        run.append(text)
    element.append(run)


def _style_paragraph(paragraph: Paragraph, style_id: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pStyle"))
    if existing is None:
        existing = OxmlElement("w:pStyle")
        p_pr.insert(0, existing)
    existing.set(qn("w:val"), style_id)


def _resolve_bound_file(raw: str, expected_sha: str, owner: Path) -> Path:
    candidate = Path(raw)
    path = candidate.resolve() if candidate.is_absolute() else (owner.parent / candidate).resolve()
    if not path.is_file() or sha256_file(path) != expected_sha.upper():
        raise DeclarativeAdapterError(f"图片证据路径或 SHA 失效: {path}")
    return path


def _image_path(item: dict[str, Any], content_plan: dict[str, Any], content_plan_path: Path) -> Path:
    decision = item.get("decision", {})
    replacement = decision.get("replacement_ref")
    if isinstance(replacement, dict) and isinstance(replacement.get("path"), str) and isinstance(replacement.get("sha256"), str):
        return _resolve_bound_file(replacement["path"], replacement["sha256"], content_plan_path)
    source_id = item.get("source_id")
    source_sha = item.get("source", {}).get("payload", {}).get("sha256")
    evidence_ids = decision.get("evidence_ids")
    matches = [
        value
        for value in content_plan.get("evidence_bindings", [])
        if isinstance(value, dict)
        and value.get("source_id") == source_id
        and isinstance(value.get("path"), str)
        and isinstance(value.get("sha256"), str)
        # 源项快照与原图证据可以并存；只选择本决策绑定的原图。
        and Path(value["path"]).suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
        and (not evidence_ids or value.get("evidence_id") in evidence_ids)
    ]
    if len(matches) != 1:
        raise DeclarativeAdapterError(f"图片槽位必须有唯一哈希绑定本地证据: {source_id}")
    path = _resolve_bound_file(matches[0]["path"], matches[0]["sha256"], content_plan_path)
    if source_sha and matches[0]["sha256"].upper() != source_sha.upper():
        # PDF 的 JP2 原图经既有预览器转 PNG 后字节不同，沿用内容校验器的像素指纹。
        visual_sha = item.get("source", {}).get("payload", {}).get("visual_sha256")
        if not visual_sha or canonical_visual_sha256(path.read_bytes()) != visual_sha.upper():
            raise DeclarativeAdapterError(f"图片证据路径或 SHA 失效: {path}")
    return path


def _text_from_item(item: dict[str, Any]) -> str | None:
    decision = item.get("decision", {})
    if isinstance(decision.get("reviewed_text"), str):
        return decision["reviewed_text"]
    payload = item.get("source", {}).get("payload", {})
    if payload.get("type") in {"text", "description"}:
        return str(payload.get("text") or payload.get("description") or "")
    return None


def _slot_material(content_plan: dict[str, Any], content_plan_path: Path) -> dict[str, list[dict[str, Any]]]:
    material: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in content_plan.get("items", []):
        if not isinstance(item, dict) or not isinstance(item.get("decision"), dict):
            raise DeclarativeAdapterError("ContentPlanV3 item/decision 无效")
        decision = item["decision"]
        if decision.get("action") in NON_RENDER_ACTIONS:
            continue
        if decision.get("action") == "human_review" or decision.get("review_required") is True:
            raise DeclarativeAdapterError(f"ContentPlanV3 仍含人工复核项: {item.get('source_id')}")
        reviewed_blocks = decision.get("reviewed_blocks")
        if isinstance(reviewed_blocks, list):
            for block in reviewed_blocks:
                if isinstance(block, dict) and isinstance(block.get("target_slot"), str):
                    material[block["target_slot"]].append({"kind": "text", "value": str(block.get("text") or ""), "source_id": item.get("source_id")})
            continue
        slot_id = decision.get("target_slot")
        if not isinstance(slot_id, str) or not slot_id:
            continue
        payload = item.get("source", {}).get("payload", {})
        if payload.get("type") == "table":
            material[slot_id].append({"kind": "table", "value": payload.get("rows", []), "source_id": item.get("source_id")})
        elif payload.get("type") == "asset":
            material[slot_id].append({"kind": "image", "value": _image_path(item, content_plan, content_plan_path), "source_id": item.get("source_id")})
        else:
            text = _text_from_item(item)
            if text is not None:
                material[slot_id].append({"kind": "text", "value": text, "source_id": item.get("source_id")})
    return material


def _product_value(source: str, content_plan: dict[str, Any]) -> str | None:
    prefix = "content_plan.product."
    if not source.startswith(prefix):
        return None
    value = content_plan.get("product", {}).get(source[len(prefix):])
    return str(value) if value is not None else None


def _text_values(values: list[dict[str, Any]], slot_id: str) -> list[str]:
    if any(value["kind"] != "text" for value in values):
        raise DeclarativeAdapterError(f"文本槽位收到非文本内容: {slot_id}")
    return [str(value["value"]) for value in values if str(value["value"])]


def _insert_table(document: Document, paragraph: Paragraph, rows: list[list[str]], slot_id: str) -> None:
    if not rows or any(not isinstance(row, list) for row in rows):
        raise DeclarativeAdapterError(f"表格槽位没有有效行: {slot_id}")
    columns = max((len(row) for row in rows), default=0)
    if columns < 1:
        raise DeclarativeAdapterError(f"表格槽位列数无效: {slot_id}")
    table = document.add_table(rows=len(rows), cols=columns)
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for row_index, row in enumerate(rows):
        # 沿用原参数表分页合同：行不跨页，首行在后续页重复。
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column_index in range(columns):
            table.cell(row_index, column_index).text = str(row[column_index]) if column_index < len(row) else ""
    paragraph._p.addnext(table._tbl)
    _replace_paragraph_text(paragraph, [])


def _insert_images(document: Document, paragraph: Paragraph, paths: list[Path], fit: str, slot_id: str,
                   display_sizes: list[tuple[float, float] | None] | None = None) -> None:
    from build_fixed_docx import image_display_size
    if fit == "cover":
        raise DeclarativeAdapterError(f"通用 adapter 无法在未知几何中安全执行 cover 裁剪: {slot_id}")
    _replace_paragraph_text(paragraph, [])
    section = document.sections[0]
    max_width = int(section.page_width - section.left_margin - section.right_margin)
    # 内联图片仍占用行基线空间；预留 12 磅，避免高图碰到页脚线。
    max_height = int(section.page_height - section.top_margin - section.bottom_margin) - 152400
    sizes = display_sizes if display_sizes is not None else [None] * len(paths)
    if len(sizes) != len(paths):
        raise DeclarativeAdapterError('图片显示尺寸与图片数量不一致')
    for path, preferred in zip(paths, sizes):
        run = paragraph.add_run()
        with Image.open(path) as image:
            width, height = image_display_size(image, max_width / 914400, max_height / 914400, preferred)
        # 原文件显示尺寸优先；缺少尺寸时按 96ppi 保守回退，超过模板范围才缩小。
        run.add_picture(str(path), width=round(width * 914400), height=round(height * 914400))


def compile_declarative_docx(
    template: Path,
    content_plan: dict[str, Any],
    content_plan_path: Path,
    slot_contract: dict[str, Any],
    program: dict[str, Any],
    output: Path,
) -> list[dict[str, Any]]:
    validate_template_program(program)
    if program.get("adapter") != ADAPTER_ID:
        raise DeclarativeAdapterError(f"adapter 不匹配: {program.get('adapter')}")
    template_facts = validate_docx_template(template)
    if template_facts["sha256"] != program.get("template_sha256") or template_facts["sha256"] != slot_contract.get("template_sha256"):
        raise DeclarativeAdapterError("模板、Program 与 SlotContract 哈希绑定不一致")
    if program.get("template_ir_sha256") != slot_contract.get("template_ir_sha256"):
        raise DeclarativeAdapterError("Program 与 SlotContract 的 Template IR 绑定不一致")
    document = Document(template)
    locations = _slot_locations(slot_contract)
    # 按原模板一次绑定锚点，避免插入表格新增段落后把后续内容写进错误单元格。
    anchors = {
        operation['slot_id']: _paragraph_for_slot(document, operation['slot_id'], locations)
        for operation in program['operations']
        if operation.get('slot_id') in locations
    }
    slots = {slot['slot_id']: slot for slot in slot_contract['slots']}
    material = _slot_material(content_plan, content_plan_path)
    display_sizes = {}
    if any(value['kind'] == 'image' for values in material.values() for value in values):
        from source_image_geometry_v3 import plan_image_display_sizes
        display_sizes = plan_image_display_sizes(content_plan, content_plan_path)
    findings: list[dict[str, Any]] = []
    for operation in program["operations"]:
        name = operation["op"]
        if name in {"retain_part", "bind_asset_role", "bind_header_footer", "section_break", "page_break", "validate_invariant"}:
            continue
        slot_id = operation.get("slot_id")
        if not isinstance(slot_id, str) or slot_id not in locations:
            raise DeclarativeAdapterError(f"DSL 操作引用未知槽位: {slot_id}")
        paragraph = anchors[slot_id]
        if name == "apply_style_ref":
            _style_paragraph(paragraph, str(operation["style_id"]))
            continue
        values = material.get(slot_id, [])
        # 仅显式声明了类型且 required=false 的图片区/表格区允许为空。
        # 有内容时仍走原类型与哈希校验；必需槽位及保守识别槽位不降级。
        expected_kind = {'insert_image': 'image', 'insert_table': 'table'}.get(name)
        if (expected_kind and not values and slots[slot_id].get('kind') == expected_kind
                and slots[slot_id].get('required') is False):
            _replace_paragraph_text(paragraph, [])
            continue
        if name == "bind_slot":
            product = _product_value(str(operation["source"]), content_plan)
            lines = [product] if product else _text_values(values, slot_id)
            if not lines and operation.get("required"):
                raise DeclarativeAdapterError(f"必需槽位没有内容: {slot_id}")
            if lines:
                _replace_paragraph_text(paragraph, lines)
        elif name == "insert_text":
            lines = _text_values(values, slot_id)
            if not lines:
                raise DeclarativeAdapterError(f"insert_text 槽位没有内容: {slot_id}")
            _replace_paragraph_text(paragraph, lines)
        elif name == "repeat_block":
            lines = _text_values(values, slot_id)
            if len(lines) < int(operation["min_items"]) or len(lines) > int(operation["max_items"]):
                raise DeclarativeAdapterError(f"repeat_block 数量越界: {slot_id}: {len(lines)}")
            _replace_paragraph_text(paragraph, lines)
        elif name == "insert_table":
            tables = [value["value"] for value in values if value["kind"] == "table"]
            if len(tables) != len(values) or len(tables) != 1:
                raise DeclarativeAdapterError(f"insert_table 要求唯一表格内容: {slot_id}")
            _insert_table(document, paragraph, tables[0], slot_id)
        elif name == "insert_image":
            images = [Path(value["value"]) for value in values if value["kind"] == "image"]
            if len(images) != len(values) or not images:
                raise DeclarativeAdapterError(f"insert_image 缺少已绑定图片: {slot_id}")
            _insert_images(document, paragraph, images, str(operation["fit"]), slot_id,
                           [display_sizes.get(value['source_id']) for value in values])
        else:
            raise DeclarativeAdapterError(f"通用 adapter 未实现 DSL 操作: {name}")
    if template_facts["sha256"] in {EMPTY_SUPPLEMENT_TEMPLATE_SHA, REPAIRED_SECOND_TEMPLATE_SHA} and not supplement_has_content(content_plan):
        text_anchor = anchors["body.supplement"]._p
        image_anchor = anchors["body.supplement.images"]._p
        heading = text_anchor.getprevious()
        # 填充完成后再省略，锚点不会漂移；保留前后正文及技术参数的必要分页。
        if (heading is None or Paragraph(heading, document._body).text != "产品资料补充"
                or text_anchor.getnext() is not image_anchor
                or any(node.xpath(".//w:drawing | .//w:pict | .//w:object | .//w:t")
                       for node in (text_anchor, image_anchor))):
            raise DeclarativeAdapterError("空补充章节边界或内容不符，拒绝省略")
        for node in (heading, text_anchor, image_anchor):
            node.getparent().remove(node)
    if template_facts["sha256"] == EMPTY_SUPPLEMENT_TEMPLATE_SHA:
        # B 线分页消融：非空补充区的高图也会挤出独立分页段；仅折叠此模板的已知边界。
        parameters = next(p for p in document.paragraphs if p.text == "五、技术参数")
        page_break = parameters._p.getnext()
        if page_break is None or not page_break.xpath('./w:r/w:br[@w:type="page"]'):
            raise DeclarativeAdapterError("技术参数分页边界不符，拒绝折叠")
        # 独立分页段在 WPS 被挤到新页后还会再翻页；段前分页不占额外行。
        parameters.paragraph_format.page_break_before = True
        page_break.getparent().remove(page_break)
    if template_facts["sha256"] == REPAIRED_SECOND_TEMPLATE_SHA:
        parameters = next(p for p in document.paragraphs if p.text == "五、技术参数")
        separator = parameters._p.getnext()
        band = separator.getnext() if separator is not None else None
        if (separator is None or separator.tag != qn("w:p")
                or separator.xpath('.//w:t | .//w:drawing | .//w:pict | .//w:br')
                or band is None or band.tag != qn("w:tbl")
                or ''.join(band.xpath('.//w:t/text()')) != '技术参数'):
            raise DeclarativeAdapterError("新版技术参数占位边界不符，拒绝压缩行高")
        # 隐藏标题及空占位段不能继承正文行高；WPS 曾因此比 Word 多留两行。
        # 保留锚点与自然分页，只约束已知的两个非正文段落。
        for node in (parameters._p, separator):
            spacing = node.get_or_add_pPr().get_or_add_spacing()
            spacing.set(qn("w:line"), "20")
            spacing.set(qn("w:lineRule"), "exact")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    validate_docx_template(output)
    return findings
